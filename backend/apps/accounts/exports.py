"""Spreadsheet export of user accounts for administrators."""

import io
from collections.abc import Iterable
from typing import Any, cast

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.accounts.models import Role, User

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = ("Correo", "Nombres", "Apellidos", "Roles", "Estado", "Fecha de registro")
DATE_FORMAT = "yyyy-mm-dd"
# Leading characters that make Excel, LibreOffice or Sheets evaluate a cell (OWASP CSV injection).
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
MIN_COLUMN_WIDTH = 12
MAX_COLUMN_WIDTH = 50
_ROLE_ORDER = {code: index for index, code in enumerate(Role.Code.values)}


def neutralize_formula(value: str) -> str:
    """An apostrophe makes spreadsheet software show the text instead of running it."""
    return f"'{value}" if value.startswith(FORMULA_PREFIXES) else value


def export_filename() -> str:
    return f"usuarios-{timezone.localdate():%Y%m%d}.xlsx"


def _role_labels(user: User) -> str:
    # Uses the prefetched roles; ordered like Role.Code, not alphabetically.
    roles = sorted(user.roles.all(), key=lambda role: _ROLE_ORDER.get(role.code, len(_ROLE_ORDER)))
    return ", ".join(role.get_code_display() for role in roles)


def _row(user: User) -> list[Any]:
    texts = [user.email, user.first_name, user.last_name, _role_labels(user)]
    return [
        *(neutralize_formula(text) for text in texts),
        "Activa" if user.is_active else "Inactiva",
        timezone.localdate(user.date_joined),
    ]


def build_users_workbook(users: Iterable[User]) -> bytes:
    """One sheet: bold frozen header, autofilter, date cells and widths sized to content."""
    workbook = Workbook()
    # A new workbook always starts with one regular worksheet.
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "Usuarios"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    widths = [len(header) for header in HEADERS]
    for user in users:
        row = _row(user)
        sheet.append(row)
        sheet.cell(row=sheet.max_row, column=len(row)).number_format = DATE_FORMAT
        widths = [
            max(width, len(value) if isinstance(value, str) else len(DATE_FORMAT))
            for width, value in zip(widths, row, strict=True)
        ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(width + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
