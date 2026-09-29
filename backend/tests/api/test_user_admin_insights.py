"""Admin panel insights: user filters, summary dashboard and Excel export.

RF-020, RNF-SEC-003, RNF-SEC-004.
"""

import io
from datetime import datetime, time, timedelta

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import AuditLog, Role, User, UserRole
from apps.accounts.services import create_user, set_roles

pytestmark = pytest.mark.django_db

USERS_URL = "/api/v1/users/"
SUMMARY_URL = "/api/v1/users/summary/"
EXPORT_URL = "/api/v1/users/export/"
PASSWORD = "Str0ng-Passw0rd!"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_HEADERS = ["Correo", "Nombres", "Apellidos", "Roles", "Estado", "Fecha de registro"]


def _user(
    email: str,
    *codes: str,
    first_name: str = "N",
    last_name: str = "N",
    is_active: bool = True,
) -> User:
    user = User.objects.create_user(
        email=email,
        password=PASSWORD,
        first_name=first_name,
        last_name=last_name,
        is_active=is_active,
    )
    for code in codes:
        UserRole.objects.create(user=user, role=Role.objects.get(code=code))
    return user


def _client_for(user: User) -> APIClient:
    client = APIClient()
    client.force_authenticate(user=user)
    return client


@pytest.fixture
def admin() -> User:
    return _user("admin@unal.edu.co", Role.Code.ADMIN, first_name="Root", last_name="Admin")


@pytest.fixture
def admin_client(admin: User) -> APIClient:
    return _client_for(admin)


@pytest.fixture
def population(admin: User) -> dict[str, User]:
    return {
        "ana": _user(
            "ana.perez@unal.edu.co",
            Role.Code.STUDENT,
            Role.Code.MONITOR,
            first_name="Ana",
            last_name="Pérez",
        ),
        "luis": _user(
            "luis.gomez@unal.edu.co",
            Role.Code.TEACHER,
            first_name="Luis",
            last_name="Gómez",
            is_active=False,
        ),
        "maria": _user(
            "maria@unal.edu.co", Role.Code.STUDENT, first_name="María", last_name="Ana Ruiz"
        ),
    }


def _emails(response_json: dict[str, object]) -> list[str]:
    results = response_json["results"]
    assert isinstance(results, list)
    return [row["email"] for row in results]


# --- Authorization ---------------------------------------------------------------------------


@pytest.mark.parametrize("url", [SUMMARY_URL, EXPORT_URL, f"{USERS_URL}?search=ana"])
def test_rnf_sec_003_insight_endpoints_return_401_without_credentials(url: str) -> None:
    assert APIClient().get(url).status_code == 401


@pytest.mark.parametrize("url", [SUMMARY_URL, EXPORT_URL, f"{USERS_URL}?role=ADMIN"])
@pytest.mark.parametrize("code", [Role.Code.STUDENT, Role.Code.MONITOR, Role.Code.TEACHER])
def test_ca_hu11_4_insight_endpoints_return_403_for_non_admins(url: str, code: str) -> None:
    client = _client_for(_user(f"{code.lower()}@unal.edu.co", code))

    assert client.get(url).status_code == 403


# --- List filters ----------------------------------------------------------------------------


def test_list_without_filters_keeps_every_user_ordered_by_email(
    admin_client: APIClient, population: dict[str, User]
) -> None:
    response = admin_client.get(USERS_URL)

    assert response.status_code == 200
    assert _emails(response.json()) == [
        "admin@unal.edu.co",
        "ana.perez@unal.edu.co",
        "luis.gomez@unal.edu.co",
        "maria@unal.edu.co",
    ]


@pytest.mark.parametrize(
    ("term", "expected"),
    [
        # First name, and last name of another user, both match case-insensitively.
        ("ANA", ["ana.perez@unal.edu.co", "maria@unal.edu.co"]),
        ("gómez", ["luis.gomez@unal.edu.co"]),
        ("maria@", ["maria@unal.edu.co"]),
        ("ana pérez", ["ana.perez@unal.edu.co"]),
        ("  luis  ", ["luis.gomez@unal.edu.co"]),
        ("nobody", []),
    ],
)
def test_search_matches_email_names_and_full_name(
    admin_client: APIClient, population: dict[str, User], term: str, expected: list[str]
) -> None:
    response = admin_client.get(USERS_URL, {"search": term})

    assert response.status_code == 200
    assert _emails(response.json()) == expected


def test_blank_search_is_ignored(admin_client: APIClient, population: dict[str, User]) -> None:
    response = admin_client.get(USERS_URL, {"search": "   "})

    assert response.json()["count"] == 4


def test_role_filter_returns_each_user_once(
    admin_client: APIClient, population: dict[str, User]
) -> None:
    response = admin_client.get(USERS_URL, {"role": "STUDENT"})

    assert response.status_code == 200
    body = response.json()
    assert body["count"] == 2
    assert _emails(body) == ["ana.perez@unal.edu.co", "maria@unal.edu.co"]
    # The serialized roles are all of the user's roles, not only the filtered one.
    assert sorted(body["results"][0]["roles"]) == ["MONITOR", "STUDENT"]


def test_role_filter_combined_with_search_does_not_duplicate_rows(
    admin_client: APIClient, population: dict[str, User]
) -> None:
    response = admin_client.get(USERS_URL, {"role": "MONITOR", "search": "ana"})

    assert _emails(response.json()) == ["ana.perez@unal.edu.co"]


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("false", ["luis.gomez@unal.edu.co"]),
        ("true", ["admin@unal.edu.co", "ana.perez@unal.edu.co", "maria@unal.edu.co"]),
        ("True", ["admin@unal.edu.co", "ana.perez@unal.edu.co", "maria@unal.edu.co"]),
    ],
)
def test_is_active_filter(
    admin_client: APIClient, population: dict[str, User], value: str, expected: list[str]
) -> None:
    response = admin_client.get(USERS_URL, {"is_active": value})

    assert response.status_code == 200
    assert _emails(response.json()) == expected


@pytest.mark.parametrize(
    ("params", "field"),
    [
        ({"role": "SUPERUSER"}, "role"),
        ({"role": "student"}, "role"),
        ({"is_active": "maybe"}, "is_active"),
        ({"is_active": "1"}, "is_active"),
    ],
)
def test_invalid_filters_return_400(
    admin_client: APIClient, params: dict[str, str], field: str
) -> None:
    response = admin_client.get(USERS_URL, params)

    assert response.status_code == 400
    assert field in response.json()


def test_filters_keep_pagination(admin_client: APIClient) -> None:
    for index in range(25):
        _user(f"student{index:02d}@unal.edu.co", Role.Code.STUDENT)

    response = admin_client.get(USERS_URL, {"role": "STUDENT"})

    body = response.json()
    assert body["count"] == 25
    assert len(body["results"]) == 20
    assert body["next"] is not None


def test_filtered_list_query_count_does_not_grow_with_users(
    admin_client: APIClient, population: dict[str, User], django_assert_max_num_queries: object
) -> None:
    for index in range(10):
        _user(f"extra{index}@unal.edu.co", Role.Code.STUDENT, Role.Code.MONITOR)

    # Auth role check + count + page + roles prefetch; never one query per user.
    with django_assert_max_num_queries(6):  # type: ignore[operator]
        response = admin_client.get(USERS_URL, {"role": "STUDENT", "search": "unal"})

    assert response.json()["count"] == 12


# --- Summary ---------------------------------------------------------------------------------


def test_summary_counts_users_by_state_role_and_recency(
    admin_client: APIClient, population: dict[str, User]
) -> None:
    old = population["maria"]
    old.date_joined = timezone.now() - timedelta(days=45)
    old.save(update_fields=["date_joined"])

    response = admin_client.get(SUMMARY_URL)

    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 4
    assert body["active"] == 3
    assert body["inactive"] == 1
    assert body["joined_last_30_days"] == 3
    assert body["by_role"] == {"STUDENT": 2, "MONITOR": 1, "TEACHER": 1, "ADMIN": 1}
    assert body["recent_activity"] == []


def test_summary_by_role_lists_every_role_even_without_users(admin_client: APIClient) -> None:
    body = admin_client.get(SUMMARY_URL).json()

    assert body["by_role"] == {"STUDENT": 0, "MONITOR": 0, "TEACHER": 0, "ADMIN": 1}


def test_summary_recent_activity_lists_ten_newest_audit_entries(
    admin_client: APIClient, admin: User
) -> None:
    target = create_user(
        actor=admin,
        email="docente@unal.edu.co",
        password=PASSWORD,
        first_name="Laura",
        last_name="Gómez",
        roles=[Role.Code.TEACHER],
    )
    for index in range(11):
        codes: list[str] = [Role.Code.STUDENT] if index % 2 == 0 else [Role.Code.TEACHER]
        set_roles(actor=admin, user=target, roles=codes)
    assert AuditLog.objects.count() == 12

    body = admin_client.get(SUMMARY_URL).json()

    activity = body["recent_activity"]
    assert len(activity) == 10
    newest = AuditLog.objects.order_by("-created_at", "-pk").first()
    assert newest is not None
    assert activity[0] == {
        "id": newest.pk,
        "action": "ROLES_CHANGED",
        "created_at": activity[0]["created_at"],
        "actor": {"id": str(admin.pk), "email": "admin@unal.edu.co", "full_name": "Root Admin"},
        "target": {
            "id": str(target.pk),
            "email": "docente@unal.edu.co",
            "full_name": "Laura Gómez",
        },
    }
    assert activity[0]["created_at"]
    stamps = [entry["created_at"] for entry in activity]
    assert stamps == sorted(stamps, reverse=True)
    assert all(entry["action"] != "USER_CREATED" for entry in activity)


def test_summary_uses_a_constant_number_of_queries(
    admin_client: APIClient, admin: User, django_assert_max_num_queries: object
) -> None:
    for index in range(5):
        create_user(
            actor=admin,
            email=f"user{index}@unal.edu.co",
            password=PASSWORD,
            first_name="U",
            last_name=str(index),
            roles=[Role.Code.STUDENT],
        )

    # Auth role check + user totals + roles + audit feed with actor/target joined.
    with django_assert_max_num_queries(5):  # type: ignore[operator]
        response = admin_client.get(SUMMARY_URL)

    assert len(response.json()["recent_activity"]) == 5


def test_summary_route_is_not_taken_for_a_user_detail(admin_client: APIClient) -> None:
    # users/<uuid>/ must never swallow the literal "summary" or "export" segments.
    assert admin_client.get(SUMMARY_URL).status_code == 200
    assert admin_client.get(EXPORT_URL).status_code == 200


# --- Export ----------------------------------------------------------------------------------


def _sheet_rows(content: bytes) -> list[list[object]]:
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    assert sheet is not None
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_export_returns_an_xlsx_attachment_with_every_user(
    admin_client: APIClient, population: dict[str, User]
) -> None:
    response = admin_client.get(EXPORT_URL)

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    today = timezone.localdate().strftime("%Y%m%d")
    assert response["Content-Disposition"] == f'attachment; filename="usuarios-{today}.xlsx"'

    rows = _sheet_rows(response.content)
    assert rows[0] == EXPORT_HEADERS
    assert len(rows) == 5
    ana = population["ana"]
    assert rows[2] == [
        "ana.perez@unal.edu.co",
        "Ana",
        "Pérez",
        "Estudiante, Monitor",
        "Activa",
        # A real date cell, so spreadsheets sort and filter it as a date.
        datetime.combine(timezone.localdate(ana.date_joined), time()),
    ]
    assert rows[3][3:5] == ["Docente", "Inactiva"]


def test_export_is_not_paginated(admin_client: APIClient) -> None:
    for index in range(25):
        _user(f"student{index:02d}@unal.edu.co", Role.Code.STUDENT)

    rows = _sheet_rows(admin_client.get(EXPORT_URL).content)

    assert len(rows) == 1 + 26


def test_export_respects_the_list_filters(
    admin_client: APIClient, population: dict[str, User]
) -> None:
    response = admin_client.get(
        EXPORT_URL, {"role": "STUDENT", "is_active": "true", "search": "ana"}
    )

    rows = _sheet_rows(response.content)
    assert [row[0] for row in rows[1:]] == ["ana.perez@unal.edu.co", "maria@unal.edu.co"]


def test_export_rejects_invalid_filters(admin_client: APIClient) -> None:
    response = admin_client.get(EXPORT_URL, {"is_active": "nope"})

    assert response.status_code == 400
    assert "is_active" in response.json()


def test_export_header_is_bold_frozen_and_filterable(admin_client: APIClient) -> None:
    response = admin_client.get(EXPORT_URL)

    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F2"
    assert sheet.column_dimensions["A"].width > len("admin@unal.edu.co")


@pytest.mark.parametrize("dangerous", ["=HYPERLINK(1)", "+1+1", "-2", "@SUM(A1)"])
def test_export_neutralizes_spreadsheet_formulas(admin_client: APIClient, dangerous: str) -> None:
    _user("formula@unal.edu.co", Role.Code.STUDENT, first_name=dangerous, last_name="Safe")

    rows = _sheet_rows(admin_client.get(EXPORT_URL, {"search": "formula"}).content)

    assert rows[1][1] == f"'{dangerous}"
    assert rows[1][2] == "Safe"
